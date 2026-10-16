"""Upload validation and text extraction with a source reference per line (file, page, paragraph).

Allowed: .txt, .md, .pdf, .docx, at most 20 MB each. File names are reduced to a safe base name and never used as a path: the file is stored under
a generated name. The extracted text is what the existing segmenter reads, and `line_map[n-1] = [page, paragraph]` says where line n came from:
    .txt .md  page null; paragraph = number of the blank-line separated block (1-based, whole file)
    .pdf      page = PDF page (1-based); paragraph = number of the blank-line separated block on that page (1-based)
    .docx     page null; paragraph = position of the line among the document paragraphs and table cells (1-based)
"""
import io
import logging
import re
from pathlib import PurePosixPath, PureWindowsPath

from . import settings
from .errors import ApiError

_CONTROL = re.compile(r"[\x00-\x1f\x7f]")
_UNSAFE = re.compile(r"[^A-Za-z0-9._ \-()]")


def safe_filename(name):
    """Base name only, no directories, no control characters, no leading dots, at most 120 characters."""
    base = PureWindowsPath(PurePosixPath(str(name or "")).name).name  # drops folders written with / or \
    base = _CONTROL.sub("", base)
    base = _UNSAFE.sub("_", base).strip(" .")
    if not base:
        base = "upload"
    stem, dot, ext = base.rpartition(".")
    if dot and stem:
        return stem[:110] + "." + ext[:8]
    return base[:120]


def check_extension(filename):
    ext = ("." + filename.rsplit(".", 1)[1].lower()) if "." in filename else ""
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise ApiError(400, "unsupported_file_type",
                       f"{filename}: file type {ext or '(none)'} is not allowed. Allowed: {', '.join(settings.ALLOWED_EXTENSIONS)}.")
    return ext


def _paragraph_map(lines, start_paragraph=1):
    """Paragraph number per line: blank lines separate paragraphs, a blank line itself belongs to the paragraph it ends."""
    out, paragraph, in_text = [], start_paragraph, False
    for line in lines:
        if line.strip():
            in_text = True
            out.append(paragraph)
        else:
            out.append(paragraph)
            if in_text:
                paragraph += 1
                in_text = False
    return out


def _decode(raw):
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def _from_text(raw):
    text = _decode(raw).replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    paragraphs = _paragraph_map(lines)
    return text, [[None, p] for p in paragraphs], None, []


def _from_pdf(raw):
    try:
        from pypdf import PdfReader
        logging.getLogger("pypdf").setLevel(logging.ERROR)
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ApiError(422, "unreadable_file", "The PDF is encrypted; remove the password and upload it again.")
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ApiError:
        raise
    except Exception as exc:
        raise ApiError(422, "unreadable_file", f"The PDF could not be read: {exc.__class__.__name__}.") from exc
    all_lines, line_map, warnings = [], [], []
    empty = 0
    for number, page_text in enumerate(pages, start=1):
        if not page_text.strip():
            empty += 1
        lines = page_text.replace("\r", "").split("\n")
        all_lines.extend(lines)
        line_map.extend([number, p] for p in _paragraph_map(lines))
    if empty:
        warnings.append(f"{empty} of {len(pages)} PDF pages have no extractable text (scanned images are not read)")
    return "\n".join(all_lines), line_map, len(pages), warnings


def _from_docx(raw):
    try:
        import docx
        document = docx.Document(io.BytesIO(raw))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells)
    except Exception as exc:
        raise ApiError(422, "unreadable_file", f"The Word file could not be read: {exc.__class__.__name__}.") from exc
    lines = [l.replace("\n", " ") for l in lines]
    return "\n".join(lines), [[None, i] for i in range(1, len(lines) + 1)], None, []


def extract(raw, ext):
    """-> {"text", "line_map", "pages", "warnings"} for the bytes of an uploaded file."""
    if ext in (".txt", ".md"):
        text, line_map, pages, warnings = _from_text(raw)
    elif ext == ".pdf":
        text, line_map, pages, warnings = _from_pdf(raw)
    elif ext == ".docx":
        text, line_map, pages, warnings = _from_docx(raw)
    else:  # pragma: no cover - guarded by check_extension
        raise ApiError(400, "unsupported_file_type", f"File type {ext} is not allowed.")
    if not text.strip():
        warnings = list(warnings) + ["the document has no extractable text"]
    return {"text": text, "line_map": line_map, "pages": pages, "warnings": warnings}


def line_map_dict(line_map):
    """{line_number: (page, paragraph)} as the segmenter expects."""
    return {i + 1: (m[0], m[1]) for i, m in enumerate(line_map)}
