"""Document ingestion: read a requirements document and return plain text.

Supported: .txt, .html/.htm, .pdf (pypdf), .docx (python-docx), .rtf (striprtf). Legacy binary .doc is NOT supported: there is no
extraction code for it, so `load_document` raises `UnsupportedFormat` with a message that says what to do (convert to .docx, .pdf or .txt).
The HTML text extraction is moved unchanged from `experiments/component1/build_srs_requirements.py`.
No models, no network, no LLM calls.
"""
import logging
from html.parser import HTMLParser
from pathlib import Path


class UnsupportedFormat(Exception):
    """The file type cannot be read by this module."""


class _HtmlText(HTMLParser):
    BLOCK = {"p", "div", "li", "tr", "br", "h1", "h2", "h3", "h4", "h5", "h6", "table", "ul", "ol", "section", "td", "th", "pre", "dd", "dt"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "head"):
            self.skip += 1
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "head"):
            self.skip = max(0, self.skip - 1)
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def _decode(raw):
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def html_to_text(path):
    parser = _HtmlText()
    parser.feed(_decode(Path(path).read_bytes()))
    return "".join(parser.parts)


def _pdf_to_text(path):
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFormat("PDF support needs the `pypdf` package (pip install -r requirements.txt).") from exc
    logging.getLogger("pypdf").setLevel(logging.ERROR)  # pypdf logs a warning per unusual font; the text is still extracted
    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages), len(pages), sum(1 for p in pages if not p.strip())


def _docx_to_text(path):
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFormat("DOCX support needs the `python-docx` package (pip install -r requirements.txt).") from exc
    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def _rtf_to_text(path):
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedFormat("RTF support needs the `striprtf` package (pip install -r requirements.txt).") from exc
    return rtf_to_text(_decode(Path(path).read_bytes()))


def load_document(path):
    """Return (text, fmt, warnings) for a document. Raises FileNotFoundError or UnsupportedFormat."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")
    suffix = path.suffix.lower()
    warnings = []
    if suffix == ".txt":
        text, fmt = path.read_bytes().decode("utf-8-sig", errors="replace"), "txt"
    elif suffix in (".html", ".htm"):
        text, fmt = html_to_text(path), "html"
    elif suffix == ".pdf":
        text, n_pages, n_empty = _pdf_to_text(path)
        fmt = "pdf"
        if n_empty:
            warnings.append(f"{n_empty} of {n_pages} PDF pages have no extractable text (scanned images are not read)")
    elif suffix == ".docx":
        text, fmt = _docx_to_text(path), "docx"
    elif suffix == ".rtf":
        text, fmt = _rtf_to_text(path), "rtf"
    elif suffix == ".doc":
        raise UnsupportedFormat(f"{path.name}: legacy binary .doc files are not supported (no extraction code). "
                                "Convert the file to .docx, .pdf or .txt and run again.")
    else:
        raise UnsupportedFormat(f"{path.name}: unsupported file type '{suffix}'. Supported: .txt, .html, .htm, .pdf, .docx, .rtf.")
    if not text.strip():
        warnings.append("the document has no extractable text")
    return text, fmt, warnings
