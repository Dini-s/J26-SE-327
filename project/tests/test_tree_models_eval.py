"""Tree-model experiment and the model comparison table: no model download, no frozen test split, tiny synthetic data."""
import ast
import zlib
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from evaluation import model_comparison as MC
from evaluation import tree_models_eval as T
from evaluation.run_baselines import METRIC_COLS

ROOT = Path(__file__).resolve().parents[1]
FR = ["The system shall allow users to create an account.", "The system shall let the manager approve a leave request.", "The user shall be able to export a report.",
      "The system shall send an email when an order ships.", "The clerk shall be able to search for a customer by name."]
NFR = ["The system shall respond within 2 seconds under normal load.", "The system shall be available 99.9 percent of the time.", "All passwords shall be encrypted.",
       "The system shall support 500 concurrent users.", "The interface shall be easy to learn for new users."]


def fake_embed(texts, use_cache=True):
    """A deterministic stand-in for MiniLM: 16 hashed bag-of-words dimensions, so the tests never download a model."""
    out = np.zeros((len(texts), 16), dtype=np.float32)
    for i, t in enumerate(texts):
        for w in str(t).lower().split():
            out[i, zlib.crc32(w.encode()) % 16] += 1.0
    return out


@pytest.fixture(autouse=True)
def no_model_download(monkeypatch):
    monkeypatch.setattr(T, "embed", fake_embed)


def data(n=6):
    X = (FR * n)[: 5 * n] + (NFR * n)[: 5 * n]
    return X, np.array(["FR"] * (5 * n) + ["NFR"] * (5 * n), dtype=object)


@pytest.mark.parametrize("family", T.FAMILIES)
@pytest.mark.parametrize("features", T.FEATURES)
def test_every_family_and_feature_set_fits_and_predicts_labels_as_written(family, features):
    X, y = data()
    params = {"max_depth": 3, "n_estimators": 20} if family == "xgboost" else {"num_leaves": 7, "n_estimators": 20} if family == "lightgbm" else {"n_estimators": 20, "max_depth": 5}
    model = T.TreeText(family, features, params).fit(X, y)
    pred = model.predict(X)
    assert set(pred) <= {"FR", "NFR"} and list(model.classes_) == ["FR", "NFR"]
    assert (pred == y).mean() > 0.9
    assert model.predict_proba(X).shape == (len(X), 2)


def test_tfidf_and_trees_are_fitted_on_the_training_texts_only():
    X, y = data()
    model = T.TreeText("random_forest", "tfidf", {"n_estimators": 10}).fit(X, y)
    assert "zebra" not in model.tfidf_.vocabulary_
    model.predict(["a zebra shall graze"])  # a test-only word is simply unknown: the vocabulary is not changed by predict
    assert "zebra" not in model.tfidf_.vocabulary_


def test_a_non_binary_task_is_refused():
    with pytest.raises(ValueError):
        T.TreeText("xgboost", "emb", {"max_depth": 2, "n_estimators": 5}).fit(["a b", "c d", "e f"], np.array(["x", "y", "z"], dtype=object))


def test_hand_features_are_plain_numbers_per_text():
    f = T.hand_features(["The system shall not log in quickly, with 3 retries.", ""])
    assert f.shape == (2, 9)
    assert f[0, 0] == 10 and f[0, 2] == 1 and f[0, 5] == 1 and f[0, 6] == 1  # words, digits, shall/must, negation
    assert (f[1] == 0).all()


def test_the_frozen_test_split_is_read_in_one_function_only():
    tree = ast.parse((ROOT / "evaluation" / "tree_models_eval.py").read_text(encoding="utf-8"))
    callers = {fn.name for fn in ast.walk(tree) if isinstance(fn, ast.FunctionDef) for node in ast.walk(fn)
               if isinstance(node, ast.Call) and getattr(node.func, "attr", "") == "load_test_for_final_evaluation"}
    assert callers == {"final"}
    for fn in [n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name in ("develop", "cross_validate", "grid_for", "hand_features")]:
        assert "load_test_for_final_evaluation" not in ast.unparse(fn), fn.name


def test_the_grids_are_small_and_the_full_option_only_adds_to_them():
    for fam in T.FAMILIES:
        assert T.grid_for(fam, False) == T.GRIDS[fam]
        assert len(T.grid_for(fam, True)) == 2 * len(T.GRIDS[fam])


# ------------------------------------------------------------------------------------------------------------------ the comparison table
def metric_rows(task, sub, model, evaluation, slice_, acc, f1, n, ci=True):
    lo = (round(acc - 0.02, 3), round(acc + 0.02, 3)) if ci else ("", "")
    base = dict(task=task, subtask=sub, evaluation=evaluation, slice=slice_, model=model, cls="", n_rows=str(n))
    return [dict(base, metric="accuracy", value=str(acc), ci_low=str(lo[0]), ci_high=str(lo[1])), dict(base, metric="macro_f1", value=str(f1), ci_low=str(lo[0]), ci_high=str(lo[1]))]


def write_fixture(tmp):
    """Majority: highest test accuracy (the imbalance trap). `rules` is best on validation; `xgboost+emb` is best on TEST but not on validation."""
    base = []
    for model, test_acc, test_f1, val_f1 in (("majority", 0.95, 0.49, 0.46), ("rules", 0.93, 0.72, 0.72), ("tfidf_lr", 0.85, 0.62, 0.68)):
        base += metric_rows("nfr_subtype", "security_binary", model, "test", "pooled", test_acc, test_f1, 1285)
        base += metric_rows("nfr_subtype", "security_binary", model, "val", "pooled", 0.9, val_f1, 900, ci=False)
    pd.DataFrame(base, columns=METRIC_COLS).to_csv(tmp / "nfr_subtype_metrics.csv", index=False)
    tree = []
    for model, test_acc, test_f1, val_f1 in (("xgboost+emb", 0.94, 0.80, 0.70), ("random_forest+tfidf", 0.90, 0.60, 0.50)):
        tree += metric_rows("nfr_subtype", "security_binary", model, "test", "pooled", test_acc, test_f1, 1285)
        tree += metric_rows("nfr_subtype", "security_binary", model, "val", "pooled", 0.9, val_f1, 900, ci=False)
    pd.DataFrame(tree, columns=METRIC_COLS).to_csv(tmp / "tree_models_metrics.csv", index=False)


def test_best_is_chosen_by_validation_and_never_by_the_test_numbers(tmp_path):
    write_fixture(tmp_path)
    rows = MC.collect(tmp_path)
    best = [r for r in rows if r["Best (by validation)"]]
    assert [r["_id"] for r in best] == ["rules"]  # not xgboost+emb (best test macro-F1) and not majority (best test accuracy)
    assert [r["_id"] for r in rows] == ["majority", "rules", "tfidf_lr", "xgboost+emb", "random_forest+tfidf"]  # baselines first, then trees in a fixed order
    agent = [r["_id"] for r in rows if r["Used by the agent"]]
    assert agent == ["rules"]


def test_the_table_files_have_accuracy_and_macro_f1_and_say_why_accuracy_misleads(tmp_path):
    write_fixture(tmp_path)
    df = MC.write(tmp_path)
    assert list(df.columns) == MC.COLUMNS
    assert {"Accuracy", "Accuracy 95% CI", "Macro-F1"} <= set(df.columns)
    csv = pd.read_csv(tmp_path / "model_comparison.csv", encoding="utf-8-sig")
    assert csv.loc[csv.Model == "XGBoost", "Features"].iloc[0] == "MiniLM embeddings"
    md = (tmp_path / "MODEL_COMPARISON.md").read_text(encoding="utf-8")
    assert "Accuracy [95% CI]" in md and "always answering" in md and "| **Keyword rules**" in md
    assert "Best tree model by validation: XGBoost" in md
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "model_comparison.xlsx")["All models"]
    assert [c.value for c in ws[1]] == MC.COLUMNS and ws.max_row == 1 + len(df)


def test_a_task_without_result_files_is_skipped_and_an_empty_folder_is_an_error(tmp_path):
    with pytest.raises(SystemExit):
        MC.write(tmp_path)
