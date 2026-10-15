"""Tests for the end-to-end requirement pipeline: ingestion, segmentation parity, duplicate detection, seeded evaluation, agent.

Run from the project directory:  pytest tests/test_pipeline.py -q
Agent tests use the models saved by `python -m evaluation.run_baselines` (git-ignored); they are skipped with a reason if those are absent.
"""
import ast
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "experiments" / "component1"))

from agents.requirement_agent import agent as agent_mod  # noqa: E402
from agents.requirement_agent import duplicates as dup  # noqa: E402
from agents.requirement_agent import features  # noqa: E402
from agents.requirement_agent import ingestion  # noqa: E402
from agents.requirement_agent import segmenter  # noqa: E402
from evaluation import duplicates_eval as de  # noqa: E402

PROCESSED = ROOT / "data" / "component1" / "processed"
RAW_SRS = ROOT / "data" / "component1" / "raw" / "documents" / "srs"
MODELS_PRESENT = all((agent_mod.MODELS_DIR / f).exists() for f in agent_mod.MODEL_FILES.values()) and (agent_mod.MODELS_DIR / "duplicates_config.json").exists()
needs_models = pytest.mark.skipif(not MODELS_PRESENT, reason="saved models missing: run evaluation.run_baselines and evaluation.duplicates_eval --task seeded")
PARITY_COLUMNS = ["srs_record_id", "source_format", "position", "line_start", "clause_label", "source_marker", "original_text", "cleaned_text",
                  "word_count", "candidate_type", "type_rule", "extraction_rule", "length_flag"]


# ---------------------------------------------------------------------------------------------------------------------------
# ingestion + segmenter parity with the existing extraction
# ---------------------------------------------------------------------------------------------------------------------------
class TestSegmenterParity:
    @pytest.fixture(scope="class")
    def expected(self):
        return pd.read_csv(PROCESSED / "srs_requirements_extracted.csv", encoding="utf-8-sig", dtype=str, keep_default_na=False)

    def test_new_path_reproduces_srs_requirements_extracted(self, expected):
        documents = [(p, p.stem) for p in sorted((PROCESSED / "srs_extracted_text").glob("*.txt"))]
        documents += [(p, p.stem) for p in sorted((ROOT / "data/component1/raw/documents/html").iterdir()) if p.suffix.lower() in (".htm", ".html")]
        assert len(documents) == 79
        total = 0
        for path, name in documents:
            text, fmt, _ = ingestion.load_document(path)
            rows, _, _ = segmenter.segment_document(name, text, fmt)
            got = pd.DataFrame(rows)[PARITY_COLUMNS].astype(str)
            want = expected[expected.source_document == name][PARITY_COLUMNS].reset_index(drop=True)
            pd.testing.assert_frame_equal(got.reset_index(drop=True), want, check_dtype=False, obj=name)
            total += len(got)
        assert total == len(expected) == 13361

    def test_original_build_script_still_produces_identical_bytes(self):
        import build_master as bm
        import build_srs_requirements as brs
        srs, counts = brs.run()
        assert bm.to_bytes(srs) == (PROCESSED / "srs_requirements_extracted.csv").read_bytes()
        assert bm.to_bytes(counts) == (bm.AUDIT / "srs_requirements_per_document_counts.csv").read_bytes()

    def test_segmenter_toy_document(self):
        text = "\n".join(["Contents", "1. Introduction ........ 3", "Page 4", "The system shall record every login attempt.", "",
                          "Note: this sentence is an example only and shall be ignored.", "", "The user may see help pages."])
        rows, dropped, _ = segmenter.segment_document("toy", text, "txt")
        assert [r["cleaned_text"] for r in rows] == ["The system shall record every login attempt."]
        assert dropped["toc_line"] == 1 and dropped["page_number"] == 1
        assert rows[0]["srs_record_id"] == "SRS:toy:0001" and rows[0]["candidate_type"] == "FR"  # "shall record" is an FR cue


class TestIngestion:
    def test_txt_and_html(self, tmp_path):
        (tmp_path / "a.txt").write_bytes("﻿The system shall log events.".encode("utf-8"))
        (tmp_path / "b.html").write_text("<html><head><title>x</title></head><body><p>The system shall log events.</p><script>var a=1;</script></body></html>", encoding="utf-8")
        t1, f1, w1 = ingestion.load_document(tmp_path / "a.txt")
        t2, f2, w2 = ingestion.load_document(tmp_path / "b.html")
        assert (f1, f2) == ("txt", "html") and "shall log events" in t1 and "shall log events" in t2 and "var a" not in t2 and "x" not in t2.split()

    def test_rtf(self, tmp_path):
        (tmp_path / "c.rtf").write_text(r"{\rtf1\ansi The system shall back up data nightly.\par Second line.}", encoding="utf-8")
        text, fmt, _ = ingestion.load_document(tmp_path / "c.rtf")
        assert fmt == "rtf" and "shall back up data nightly" in text

    def test_docx(self, tmp_path):
        import docx
        d = docx.Document()
        d.add_paragraph("The system shall export reports as CSV.")
        t = d.add_table(rows=1, cols=1)
        t.rows[0].cells[0].text = "The system shall list open orders."
        d.save(str(tmp_path / "d.docx"))
        text, fmt, _ = ingestion.load_document(tmp_path / "d.docx")
        assert fmt == "docx" and "export reports as CSV" in text and "list open orders" in text

    def test_pdf(self):
        pdfs = sorted(RAW_SRS.glob("*.pdf"), key=lambda p: p.stat().st_size)
        text, fmt, _ = ingestion.load_document(pdfs[0])
        assert fmt == "pdf" and len(text.split()) > 100

    def test_legacy_doc_is_reported_unsupported(self, tmp_path):
        (tmp_path / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0")
        with pytest.raises(ingestion.UnsupportedFormat, match="legacy binary .doc"):
            ingestion.load_document(tmp_path / "old.doc")

    def test_unknown_missing_and_empty(self, tmp_path):
        (tmp_path / "x.xyz").write_text("hi")
        with pytest.raises(ingestion.UnsupportedFormat):
            ingestion.load_document(tmp_path / "x.xyz")
        with pytest.raises(FileNotFoundError):
            ingestion.load_document(tmp_path / "missing.txt")
        (tmp_path / "e.txt").write_text("   ")
        assert "no extractable text" in ingestion.load_document(tmp_path / "e.txt")[2][0]

    def test_raw_srs_format_counts(self):
        suffixes = pd.Series([p.suffix.lower() for p in RAW_SRS.iterdir()]).value_counts().to_dict()
        assert suffixes == {".pdf": 62, ".doc": 13, ".rtf": 1}


class TestNormalisation:
    def test_matches_processed_normalised_text(self):
        import build_master as bm
        m = pd.read_csv(PROCESSED / "master_requirements.csv", encoding="utf-8-sig", dtype=str, keep_default_na=False).iloc[::7]
        for original, normalised in zip(m.requirement_text_original, m.requirement_text_normalized):
            assert features.normalize_for_models(original) == normalised == bm.normalize_text(original)


# ---------------------------------------------------------------------------------------------------------------------------
# duplicates
# ---------------------------------------------------------------------------------------------------------------------------
def _unit(*v):
    v = np.array(v, dtype=np.float32)
    return v / np.linalg.norm(v)


class TestDuplicates:
    emb = np.vstack([_unit(1, 0, 0), _unit(1, 0.05, 0), _unit(0, 1, 0), _unit(0.05, 1, 0), _unit(0, 0, 1), _unit(0.02, 1, 0.01)])

    def test_groups_and_singletons(self):
        ids = dup.group_ids(dup.cluster_labels(self.emb, 0.95))
        assert ids[0] == ids[1] != "" and ids[2] == ids[3] == ids[5] != "" and ids[4] == "" and ids[0] != ids[2]
        assert ids[0] == "DUP0001" and ids[2] == "DUP0002"

    def test_deterministic(self):
        assert dup.group_ids(dup.cluster_labels(self.emb, 0.95)) == dup.group_ids(dup.cluster_labels(self.emb, 0.95))

    def test_similarity_pairs_range(self):
        i, j, s = dup.similarity_pairs(self.emb, 0.95)
        assert set(zip(i.tolist(), j.tolist())) == {(0, 1), (2, 3), (2, 5), (3, 5)}
        i2, j2, s2 = dup.similarity_pairs(self.emb, 0.95, high=0.9999)
        assert all(x < 0.9999 for x in s2)

    def test_threshold_controls_grouping(self):
        assert set(dup.group_ids(dup.cluster_labels(self.emb, 0.9999))) == {""}

    def test_exact_key_and_single_row(self):
        assert dup.exact_key("The  System, SHALL!") == dup.exact_key("the system shall")
        assert dup.group_ids(dup.cluster_labels(self.emb[:1], 0.9)) == [""]

    def test_missing_config_message(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="duplicates_eval"):
            dup.load_threshold(tmp_path / "nope.json")


class TestSeededEvaluation:
    text = "The system shall allow the user to quickly export all reports, when the report is complete."

    def test_each_edit_changes_the_text_and_stays_close(self):
        for name, fn in de.EDITS.items():
            new = fn(self.text)
            assert new and new != self.text, name
            assert de.jaccard(self.text, new) >= 0.5, name

    def test_make_positive_is_deterministic(self):
        assert de.make_positive(self.text) == de.make_positive(self.text)

    def test_hard_negative_shares_template_but_not_object(self):
        a = "The system shall display the customer name on the invoice."
        b = "The system shall display the order number on the report."
        prefix = {" ".join(w.lower() for w in t.split()[:3]): [a, b] for t in (a, b)}
        new = de.make_hard_negative(a, prefix)
        assert new and new != a and new.startswith("The system shall display the customer name on the") and de.jaccard(a, new) >= 0.6
        assert de.make_hard_negative(a, {}) is None

    def test_build_pairs_labels_and_disjoint_bases(self):
        texts = [f"The system shall {v} the {o} for the registered user account when requested." for v in ("display", "print", "export", "delete", "update", "create")
                 for o in ("invoice", "report", "profile", "order", "ticket", "schedule", "balance")]
        df = de.build_pairs(texts, 5, "toy")
        assert set(df.kind) <= {"positive", "hard_negative", "easy_negative"} and df.seeded_eval_only.all()
        assert (df[df.kind == "positive"].label == "duplicate").all() and (df[df.kind != "positive"].label == "not_duplicate").all()
        assert len(df.pair_id) == len(set(df.pair_id))

    def test_saved_seeded_files_are_labelled_and_use_no_test_data(self):
        out = ROOT / "evaluation" / "component1" / "results" / "duplicates_eval"
        if not (out / "seeded_pairs_eval.csv").exists():
            pytest.skip("seeded evaluation not run yet")
        for f in ("seeded_pairs_dev.csv", "seeded_pairs_eval.csv"):
            d = pd.read_csv(out / f, dtype=str, encoding="utf-8-sig", keep_default_na=False)
            assert (d.seeded_eval_only == "True").all()
        src = (ROOT / "evaluation" / "duplicates_eval.py").read_text(encoding="utf-8")
        assert "load_test_for_final_evaluation" not in src and '("test")' not in src
        assert 'requirement_texts("train")' in src and 'requirement_texts("val")' in src

    def test_config_threshold_matches_a_sweep_value(self):
        cfg = ROOT / "models" / "component1" / "duplicates_config.json"
        if not cfg.exists():
            pytest.skip("config not created yet")
        sweep = pd.read_csv(ROOT / "evaluation/component1/results/duplicates_eval/threshold_sweep_dev.csv")
        best = sweep.sort_values(["f1", "threshold"], ascending=[False, False]).iloc[0]
        assert abs(dup.load_threshold(cfg) - best.threshold) < 1e-9


class TestReviewFile:
    def test_shape_and_rules(self):
        path = ROOT / "evaluation" / "component1" / "results" / "duplicates_review_top50.csv"
        if not path.exists():
            pytest.skip("review file not built")
        d = pd.read_csv(path, dtype=str, encoding="utf-8-sig", keep_default_na=False)
        assert len(d) == 50 and (d.selection == "top25_highest_cosine").sum() == 25 and (d.selection == "band_0.80_0.90_sampled").sum() == 25
        assert set(d.scope) <= {"within_document", "across_documents"} and (d.is_duplicate == "").all() and (d.reviewer_note == "").all()
        band = d[d.selection == "band_0.80_0.90_sampled"].cosine.astype(float)
        assert band.between(0.80, 0.90, inclusive="left").all()
        assert all(dup.exact_key(a) != dup.exact_key(b) for a, b in zip(d.text_a, d.text_b))
        assert d.pair_no.astype(int).tolist() == list(range(1, 51))
        assert (d[(d.document_a == "1999 - tcs") | (d.document_b == "1999 - tcs")].tcs_flag == "tcs_hold").all()


# ---------------------------------------------------------------------------------------------------------------------------
# agent
# ---------------------------------------------------------------------------------------------------------------------------
DOC = "\n".join([
    "1. The system shall encrypt all stored passwords using a strong algorithm.",
    "2. The system shall display the list of open orders to the clerk.",
    "3. The response time of the search function should be adequate for most users.",
    "4. The system shall display the list of open orders to the clerk.",
    "5. The system shall allow an administrator to reset user passwords."])


@needs_models
class TestAgent:
    @pytest.fixture(scope="class")
    def result(self, tmp_path_factory):
        path = tmp_path_factory.mktemp("doc") / "toy requirements.txt"
        path.write_text(DOC, encoding="utf-8")
        return agent_mod.run(path), path

    def test_columns_and_rows(self, result):
        df, _ = result
        assert list(df.columns) == agent_mod.COLUMNS and len(df) == 5
        assert set(df.requirement_type) <= {"FR", "NFR"} and df.ambiguity_score.between(0, 1).all() and df.nfr_probability.between(0, 1).all()
        assert (df.boundary_status == "auto_unverified").all()

    def test_explanations_are_in_the_text(self, result):
        df, _ = result
        import re
        for r in df.itertuples():
            for w in filter(None, r.weak_words_matched.split(";")):
                assert re.search(rf"(?<![a-z0-9]){re.escape(w)}(?![a-z0-9])", r.text, re.I), (w, r.text)
            for w in filter(None, r.security_cue_words.split(";")):
                assert w in r.text.lower()
            assert r.is_security == bool(r.security_cue_words)

    def test_security_and_duplicates(self, result):
        df, _ = result
        assert df.is_security.iloc[0] and df.is_security.iloc[4]
        g = df.duplicate_group_id
        assert g.iloc[1] == g.iloc[3] != "" and g.iloc[0] == ""

    def test_deterministic(self, result):
        df, path = result
        pd.testing.assert_frame_equal(df, agent_mod.run(path))

    def test_cli_writes_csv(self, result, tmp_path):
        _, path = result
        out = tmp_path / "o.csv"
        assert agent_mod.main([str(path), "--out", str(out)]) == 0
        assert len(pd.read_csv(out, encoding="utf-8-sig")) == 5

    def test_missing_model_stops_with_clear_message(self, result, tmp_path, capsys):
        _, path = result
        with pytest.raises(agent_mod.ModelMissing, match="fr_nfr_fr_nfr_emb_lr.joblib"):
            agent_mod.run(path, models_dir=tmp_path)
        assert agent_mod.main([str(path), "--out", str(tmp_path / "x.csv"), "--models-dir", str(tmp_path)]) == 2
        err = capsys.readouterr().err
        assert "STOP" in err and "run_baselines" in err

    def test_missing_duplicate_config_stops(self, result, tmp_path):
        _, path = result
        for f in agent_mod.MODEL_FILES.values():
            shutil.copy(agent_mod.MODELS_DIR / f, tmp_path / f)
        with pytest.raises(agent_mod.ModelMissing, match="duplicates_config.json"):
            agent_mod.run(path, models_dir=tmp_path)

    def test_unsupported_document(self, tmp_path, capsys):
        (tmp_path / "x.doc").write_bytes(b"\xd0\xcf")
        assert agent_mod.main([str(tmp_path / "x.doc"), "--out", str(tmp_path / "o.csv")]) == 2


class TestNoLlmNoNetwork:
    def test_new_modules_import_no_llm_or_network_library(self):
        banned = {"anthropic", "openai", "requests", "urllib", "http", "socket", "httpx", "langchain", "cohere", "google"}
        for rel in ("agents/requirement_agent/agent.py", "agents/requirement_agent/ingestion.py", "agents/requirement_agent/segmenter.py",
                    "agents/requirement_agent/duplicates.py", "evaluation/duplicates_eval.py"):
            tree = ast.parse((ROOT / rel).read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                names = [a.name for a in node.names] if isinstance(node, ast.Import) else [node.module or ""] if isinstance(node, ast.ImportFrom) else []
                for n in names:
                    assert n.split(".")[0] not in banned, f"{rel} imports {n}"

    def test_requirements_pin_new_dependencies(self):
        req = (ROOT / "requirements.txt").read_text(encoding="utf-8").split()
        for pin in ("pypdf==6.20.0", "python-docx==1.2.0", "striprtf==0.0.33", "joblib==1.6.0", "sentence-transformers==6.1.0"):
            assert pin in req


class TestExtractionPrecision:
    def test_not_available_when_nothing_reviewed(self):
        from evaluation import extraction_precision as ep
        df = pd.DataFrame({"is_requirement": ["", ""], "source_document": ["a", "b"], "extraction_rule": ["modal_sentence"] * 2})
        text, result = ep.summarise(df)
        assert result is None and "NOT AVAILABLE: 0 of 2" in text

    def test_precision_and_ci_on_filled_toy_file(self):
        from evaluation import extraction_precision as ep
        df = pd.DataFrame({"is_requirement": ["yes"] * 8 + ["no"] * 2 + ["", "maybe"], "source_document": list("abcdefghijkl"),
                           "extraction_rule": ["modal_sentence"] * 12})
        text, (p, lo, hi) = ep.summarise(df)
        assert p == 0.8 and lo <= p <= hi and "Reviewed rows: 10 of 12" in text
        assert ep.precision_with_ci([1, 1, 0, 0]) == ep.precision_with_ci([1, 1, 0, 0])

    def test_real_spot_check_file_has_the_review_columns(self):
        d = pd.read_csv(PROCESSED / "tasks" / "srs_extraction_spot_check.csv", encoding="utf-8-sig", dtype=str, keep_default_na=False)
        assert len(d) == 100 and "is_requirement" in d.columns and "reviewer_note" in d.columns
