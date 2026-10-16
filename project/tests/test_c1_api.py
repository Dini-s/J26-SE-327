"""Component 1 API: uploads, Part A analysis, filters, generation (mock/replay/unreachable), caching, decisions, export, error shape.

Each test gets its own workspace in a temporary folder and its own embedding cache, so neither outputs/ nor models/ is written. Generation uses a
scripted backend unless the test is about the replay or Ollama backend. Tests that need the saved Part A models are skipped if they are missing.
"""
import csv
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from agents.requirement_agent import agent as agent_mod
from agents.requirement_agent import features
from agents.requirement_agent.api import settings
from agents.requirement_agent.api.app import create_app
from agents.requirement_agent.api.gen_worker import mock_reply
from agents.requirement_agent.generation import MockBackend

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = (ROOT / "tests" / "fixtures" / "sample_srs.txt").read_bytes()
MODELS_PRESENT = all((agent_mod.MODELS_DIR / f).exists() for f in agent_mod.MODEL_FILES.values()) and (agent_mod.MODELS_DIR / "duplicates_config.json").exists()
pytestmark = pytest.mark.skipif(not MODELS_PRESENT, reason="saved models missing in models/component1/part_a_active")

CONFLICT_DOC = (b"1. Performance\n\n1.1 The system shall return search results within 2 seconds.\n1.2 The system shall return search results within 5 seconds.\n"
                b"1.3 The system shall allow users to export reports as PDF.\n1.4 The system shall not allow users to export reports as PDF.\n")


class CountingBackend(MockBackend):
    def __init__(self):
        super().__init__(mock_reply)


@pytest.fixture
def make_client(tmp_path, monkeypatch):
    monkeypatch.setattr(features, "CACHE_DIR", features.CACHE_DIR)  # registered so the app's cache redirect is undone after the test
    monkeypatch.setattr(features, "CACHE_FILE", features.CACHE_FILE)
    monkeypatch.setattr(features, "_CACHE", features._CACHE)
    clients = []

    def build(**kw):
        kw.setdefault("workspace", tmp_path / "ws")
        kw.setdefault("backend_name", "mock")
        if "backend" not in kw and kw["backend_name"] == "mock":
            kw["backend"] = CountingBackend()
        c = TestClient(create_app(**kw))
        c.__enter__()
        clients.append(c)
        return c

    yield build
    for c in clients:
        c.__exit__(None, None, None)


@pytest.fixture
def client(make_client):
    return make_client()


def upload(c, files, pid=1):
    return c.post(f"/api/projects/{pid}/sources", files=[("files", f) for f in files])


def analyze(c, files=(("sample_srs.txt", SAMPLE, "text/plain"),), pid=1):
    if files:
        r = upload(c, files, pid)
        assert r.status_code == 201, r.text
    j = c.post(f"/api/projects/{pid}/analyze", params={"policy": "all"})
    assert j.status_code == 202, j.text
    assert c.app.state.manager.wait_idle(120)
    return j.json()


def reqs(c, pid=1, **params):
    return c.get(f"/api/projects/{pid}/requirements", params=params).json()


def error_of(r):
    body = r.json()
    assert set(body) == {"error"} and set(body["error"]) == {"code", "message"}, body
    return body["error"]


def make_pdf(pages):
    """A tiny valid PDF (Helvetica text) built by hand; pages is a list of line lists."""
    esc = lambda s: s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    n = len(pages)
    objs = {1: "<< /Type /Catalog /Pages 2 0 R >>", 2: "<< /Type /Pages /Kids [" + " ".join(f"{4 + 2 * i} 0 R" for i in range(n)) + f"] /Count {n} >>",
            3: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"}
    for i, lines in enumerate(pages):
        body = "BT /F1 11 Tf 14 TL 50 750 Td " + " T* ".join(f"({esc(l)}) Tj" for l in lines) + " ET"
        objs[4 + 2 * i] = f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        objs[5 + 2 * i] = f"<< /Length {len(body)} >>\nstream\n{body}\nendstream"
    out, offsets = bytearray(b"%PDF-1.4\n"), {}
    for k in sorted(objs):
        offsets[k] = len(out)
        out += f"{k} 0 obj\n{objs[k]}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for k in sorted(objs):
        out += f"{offsets[k]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


# ------------------------------------------------------------------------------------------------------------------ health, projects
def test_health_and_default_project(client):
    h = client.get("/api/health").json()
    assert h == {"status": "ok", "part_a_loaded": True, "generation_backend": "mock", "generation_reachable": True, "model": "mock"}
    projects = client.get("/api/projects").json()
    assert [p["name"] for p in projects] == ["Core Banking Platform"] and projects[0]["id"] == 1
    p = client.post("/api/projects", json={"name": "  Second project "}).json()
    assert p["name"] == "Second project" and p["sources"] == 0
    assert len(client.get("/api/projects").json()) == 2


def test_error_shape_for_unknown_project_and_bad_input(client):
    assert error_of(client.get("/api/projects/99/summary"))["code"] == "project_not_found"
    assert error_of(client.get("/api/requirements/REQ-999"))["code"] == "requirement_not_found"
    assert error_of(client.get("/api/jobs/nope"))["code"] == "job_not_found"
    r = client.post("/api/projects", json={"name": ""})
    assert r.status_code == 422 and error_of(r)["code"] == "validation_error"
    assert error_of(client.get("/api/nothing"))["code"] == "not_found"
    assert error_of(client.post("/api/projects/1/analyze"))["code"] == "no_sources"


# ------------------------------------------------------------------------------------------------------------------------ uploads
def test_upload_txt_and_list_sources(client):
    r = upload(client, [("spec.txt", SAMPLE, "text/plain")])
    assert r.status_code == 201
    body = r.json()
    assert [s["filename"] for s in body["accepted"]] == ["spec.txt"] and body["rejected"] == []
    listed = client.get("/api/projects/1/sources").json()
    assert listed[0]["size_bytes"] == len(SAMPLE) and listed[0]["pages"] is None


def test_upload_docx_pdf_and_md_keep_page_and_paragraph(client):
    import docx
    d = docx.Document()
    d.add_paragraph("Introduction paragraph.")
    d.add_paragraph("3.1 The system shall lock a user account after five failed login attempts.")
    buf = io.BytesIO()
    d.save(buf)
    pdf = make_pdf([["1 Overview", "1.1 The system shall display the account balance to the customer."], ["2 Security", "2.1 The system shall encrypt all stored passwords."]])
    md = b"# Spec\n\n- The system shall export reports as PDF.\n"
    r = upload(client, [("a.docx", buf.getvalue(), "application/octet-stream"), ("b.pdf", pdf, "application/pdf"), ("c.md", md, "text/markdown")])
    assert r.status_code == 201 and len(r.json()["accepted"]) == 3, r.text
    assert [s["pages"] for s in r.json()["accepted"]] == [None, 2, None]
    client.post("/api/projects/1/analyze")
    assert client.app.state.manager.wait_idle(120)
    items = {i["source_file"]: i for i in reqs(client)["items"]}
    assert set(items) == {"a.docx", "b.pdf", "c.md"}
    pdf_rows = [i for i in reqs(client)["items"] if i["source_file"] == "b.pdf"]
    assert sorted(i["page"] for i in pdf_rows) == [1, 2]
    assert items["a.docx"]["page"] is None and items["a.docx"]["paragraph"] == 2
    assert items["c.md"]["page"] is None and items["c.md"]["paragraph"] >= 1


def test_bad_file_type_is_rejected(client):
    r = upload(client, [("virus.exe", b"MZ", "application/octet-stream")])
    assert r.status_code == 400 and error_of(r)["code"] == "unsupported_file_type"
    assert client.get("/api/projects/1/sources").json() == []


def test_mixed_upload_accepts_the_good_file_and_reports_the_bad_one(client):
    r = upload(client, [("good.txt", SAMPLE, "text/plain"), ("bad.html", b"<p>x</p>", "text/html")])
    assert r.status_code == 201
    assert [f["filename"] for f in r.json()["accepted"]] == ["good.txt"]
    assert r.json()["rejected"][0]["code"] == "unsupported_file_type"


def test_oversize_file_is_rejected(client, monkeypatch):
    monkeypatch.setattr(settings, "MAX_UPLOAD_BYTES", 100)
    r = upload(client, [("big.txt", b"x" * 500, "text/plain")])
    assert r.status_code == 413 and error_of(r)["code"] == "file_too_large"


def test_empty_and_corrupt_files_are_rejected(client):
    assert error_of(upload(client, [("empty.txt", b"", "text/plain")]))["code"] == "empty_file"
    r = upload(client, [("broken.pdf", b"not a pdf at all", "application/pdf")])
    assert r.status_code == 422 and error_of(r)["code"] == "unreadable_file"
    r = upload(client, [("broken.docx", b"not a zip", "application/octet-stream")])
    assert r.status_code == 422 and error_of(r)["code"] == "unreadable_file"


def test_path_traversal_file_name_is_neutralised(client, tmp_path):
    r = upload(client, [("../../../evil.txt", SAMPLE, "text/plain")])
    assert r.status_code == 201 and r.json()["accepted"][0]["filename"] == "evil.txt"
    stored = list((tmp_path / "ws" / "uploads").rglob("*.txt"))
    assert len(stored) == 1 and (tmp_path / "ws" / "uploads") in stored[0].parents and stored[0].name != "evil.txt"
    assert not (tmp_path / "evil.txt").exists() and not (tmp_path.parent / "evil.txt").exists()


# ----------------------------------------------------------------------------------------------------------------- analysis, summary
def test_analyze_summary_and_requirement_detail(client):
    job = analyze(client)
    assert job["part_a_done"] is True and job["requirements"] == 3
    j = client.get(f"/api/jobs/{job['job_id']}").json()
    assert j["stage"] == "share" and j["generation"] == {"total": 3, "done": 3, "failed": 0, "queued": 0, "running": 0}

    s = client.get("/api/projects/1/summary").json()
    assert s["requirements"] == 3 and s["duplicates"] == 1 and s["duplicate_groups"] == 1
    assert s["ambiguities"]["count"] == 1 and (s["task_drafts"]["generated"], s["task_drafts"]["total"]) == (3, 3)
    assert s["conflicts"] == 0 and s["open_findings"] >= 3

    items = reqs(client)["items"]
    assert [i["id"] for i in items] == ["REQ-001", "REQ-002", "REQ-003"] and items[0]["uid"] == "p1-REQ-001"
    d = client.get("/api/requirements/REQ-003").json()
    assert d["text"].startswith("The system shall be easy to use") and d["ambiguity"]["label"] == "defect"
    assert d["ambiguity"]["weak_words"] == ["appropriate"]
    span = d["ambiguity"]["weak_word_spans"][0]
    assert d["text"][span["start"]:span["end"]].lower() == "appropriate"
    assert "appropriate" in d["explanation"] and "probability" in d["explanation"]
    assert d["risk"]["level"] in ("MEDIUM", "HIGH") and d["generation"]["state"] == "done" and len(d["generation"]["acceptance_criteria"]) == 2
    first = client.get("/api/requirements/REQ-001").json()
    assert first["duplicates"][0]["id"] == "REQ-002" and first["duplicates"][0]["similarity"] >= 0.955
    assert "REQ-002" in first["explanation"] and first["source"]["file"] == "sample_srs.txt" and first["source"]["paragraph"] == 2


def test_candidate_conflicts_are_found_and_labelled(client):
    analyze(client, files=(("conflicts.txt", CONFLICT_DOC, "text/plain"),))
    s = client.get("/api/projects/1/summary").json()
    assert s["conflicts"] >= 2
    conflicting = reqs(client, type="conflict")["items"]
    assert len(conflicting) >= 4
    d = client.get("/api/requirements/REQ-001").json()
    assert d["conflicts"] and "needs human decision" in d["conflicts"][0]["reason"] and d["conflicts"][0]["id"] == "REQ-002"
    assert "Candidate conflict, needs human decision" in d["explanation"]


def test_filters_search_and_paging(client):
    analyze(client)
    assert reqs(client)["total"] == 3
    assert [i["id"] for i in reqs(client, type="ambiguity")["items"]] == ["REQ-003"]
    assert sorted(i["id"] for i in reqs(client, type="duplicate")["items"]) == ["REQ-001", "REQ-002"]
    assert reqs(client, type="conflict")["total"] == 0
    assert [i["id"] for i in reqs(client, type="security")["items"]] == ["REQ-001", "REQ-002"]
    level = client.get("/api/requirements/REQ-003").json()["risk"]["level"]
    assert "REQ-003" in [i["id"] for i in reqs(client, risk=level.lower())["items"]]
    assert reqs(client, risk="HIGH")["total"] + reqs(client, risk="MEDIUM")["total"] + reqs(client, risk="LOW")["total"] + reqs(client, risk="NONE")["total"] == 3
    assert [i["id"] for i in reqs(client, q="easy to use")["items"]] == ["REQ-003"]
    assert [i["id"] for i in reqs(client, q="req-002")["items"]] == ["REQ-002"]
    assert reqs(client, q="no such words")["total"] == 0
    page = reqs(client, limit=2, offset=1)
    assert page["total"] == 3 and [i["id"] for i in page["items"]] == ["REQ-002", "REQ-003"]
    assert error_of(client.get("/api/projects/1/requirements", params={"risk": "EXTREME"}))["code"] == "invalid_parameter"
    assert error_of(client.get("/api/projects/1/requirements", params={"type": "other"}))["code"] == "invalid_parameter"
    assert error_of(client.get("/api/projects/1/requirements", params={"limit": 0}))["code"] == "invalid_parameter"


# ------------------------------------------------------------------------------------------------------------------------ generation
def test_single_generate_and_cache(make_client):
    backend = CountingBackend()
    c = make_client(backend=backend)
    analyze(c)
    assert len(backend.calls) == 3
    analyze(c, files=())  # analyse again with the same sources: ids stay, results come from the cache
    assert len(backend.calls) == 3
    assert c.get("/api/requirements/REQ-001").json()["generation"]["message"] == "from cache"
    r = c.post("/api/requirements/REQ-001/generate")
    assert r.status_code == 202 and r.json() == {"uid": "p1-REQ-001", "generation_state": "queued"}
    assert c.app.state.manager.wait_idle(60)
    assert len(backend.calls) == 4  # an explicit regenerate skips the cache
    g = c.get("/api/requirements/REQ-001").json()["generation"]
    assert g["state"] == "done" and g["status"] == "ok" and g["user_story"].startswith("As a") and g["attempts"] == 1


def test_invalid_replies_end_in_manual_review(make_client):
    c = make_client(backend=MockBackend(["not json"]))
    analyze(c)
    g = c.get("/api/requirements/REQ-001").json()["generation"]
    assert g["state"] == "manual_review" and g["manual_review"] is True and g["attempts"] == 3 and g["raw_output"] == "not json"
    assert reqs(c)["items"][0]["generation_state"] == "manual_review"


def test_needs_clarification_state(make_client):
    reply = json.dumps({"status": "needs_clarification", "clarification_questions": ["What does fast mean?"], "rationale": "Not measurable."})
    c = make_client(backend=MockBackend([reply]))
    analyze(c)
    g = c.get("/api/requirements/REQ-003").json()["generation"]
    assert g["state"] == "needs_clarification" and g["clarification_questions"] == ["What does fast mean?"] and g["user_story"] == ""
    assert c.get("/api/projects/1/summary").json()["task_drafts"]["generated"] == 0


def test_replay_backend_uses_stored_outputs_and_never_invents_missing_ones(make_client, tmp_path):
    texts = ["The system shall lock a user account after five failed login attempts."]
    out = {"status": "ok", "user_story": "As a security officer, I want accounts locked, so that guessing is slowed.", "technical_task": "Lock after five failures.",
           "acceptance_criteria": ["Given five failures, when the fifth fails, then the account is locked.", "Given a correct login, when it is entered, then access is granted."],
           "rationale": "Role assumed."}
    replay = tmp_path / "replay.json"
    replay.write_text(json.dumps([{"id": "X:1", "requirement_text": texts[0], "output": json.dumps(out)}]), encoding="utf-8")
    c = make_client(backend=None, backend_name="replay", replay_file=replay)
    assert c.get("/api/health").json()["generation_backend"] == "replay"
    analyze(c)
    hit = c.get("/api/requirements/REQ-001").json()["generation"]
    assert hit["state"] == "done" and hit["user_story"] == out["user_story"]
    miss = c.get("/api/requirements/REQ-003").json()["generation"]
    assert miss["state"] == "failed" and miss["message"] == "not in replay file" and miss["user_story"] == "" and miss["acceptance_criteria"] == []
    assert c.get("/api/projects/1/summary").json()["task_drafts"]["generated"] == 1  # REQ-002 differs from the stored text, so it is not guessed


def test_missing_replay_file_does_not_crash_the_server(make_client, tmp_path):
    c = make_client(backend=None, backend_name="replay", replay_file=tmp_path / "missing.json")
    assert c.get("/api/health").json()["generation_reachable"] is False
    analyze(c)
    g = c.get("/api/requirements/REQ-001").json()["generation"]
    assert g["state"] == "failed" and "replay file not found" in g["message"]


def test_ollama_unreachable_marks_failed_and_keeps_serving(make_client, monkeypatch):
    monkeypatch.setenv("LLM_BASE_URL", "http://127.0.0.1:9/v1")
    monkeypatch.setenv("LLM_MODEL", "test-model")
    c = make_client(backend=None, backend_name="ollama")
    h = c.get("/api/health").json()
    assert h["generation_backend"] == "ollama" and h["generation_reachable"] is False and h["model"] == "test-model" and h["part_a_loaded"] is True
    analyze(c)
    states = [i["generation_state"] for i in reqs(c)["items"]]
    assert states == ["failed"] * 3
    g = c.get("/api/requirements/REQ-001").json()["generation"]
    assert g["message"].startswith("AI engine unavailable: ") and "could not get a reply" in g["message"]
    assert g["backend"] == "" and g["model"] == ""  # nothing was generated, so no provenance is claimed
    assert c.get("/api/health").status_code == 200 and c.get("/api/projects/1/summary").status_code == 200  # still serving
    j = c.get("/api/projects/1/summary").json()
    assert j["requirements"] == 3 and j["task_drafts"]["generated"] == 0


def test_each_draft_says_which_backend_and_model_wrote_it(make_client):
    c = make_client(backend=CountingBackend(), model_name="demo-model")
    analyze(c)
    g = c.get("/api/requirements/REQ-001").json()["generation"]
    assert g["state"] == "done" and g["backend"] == "mock" and g["model"] == "demo-model"
    analyze(c, files=())  # second run is served from the cache and still says where the text came from
    g = c.get("/api/requirements/REQ-001").json()["generation"]
    assert g["message"] == "from cache" and g["backend"] == "mock" and g["model"] == "demo-model"
    assert c.get("/api/requirements/REQ-001").json()["generation"]["raw_output"] == ""


def test_ollama_health_needs_the_configured_model_to_be_served(make_client, monkeypatch):
    import urllib.request

    class Reply:
        def __init__(self, ids):
            self.ids = ids

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self):
            return json.dumps({"object": "list", "data": [{"id": i} for i in self.ids]}).encode()

    monkeypatch.setenv("LLM_BASE_URL", "http://127.0.0.1:9/v1")
    monkeypatch.setenv("LLM_MODEL", "c1-req")
    c = make_client(backend=None, backend_name="ollama")
    for served, expected in ([["c1-req:latest"], True], [["c1-req"], True], [["llama3:latest"], False], [[], False]):
        monkeypatch.setattr(urllib.request, "urlopen", lambda *a, _s=served, **k: Reply(_s))
        assert c.get("/api/health").json()["generation_reachable"] is expected, served


def test_a_failed_draft_can_be_retried_once_the_engine_is_back(make_client):
    class Flaky(MockBackend):
        up = False

        def generate(self, system, user):
            if not self.up:
                from agents.requirement_agent.generation import BackendError
                raise BackendError("connection refused")
            return super().generate(system, user)

    b = Flaky(mock_reply)
    c = make_client(backend=b)
    analyze(c)
    g = c.get("/api/requirements/REQ-001").json()["generation"]
    assert g["state"] == "failed" and g["message"].startswith("AI engine unavailable: ")
    b.up = True
    assert c.post("/api/requirements/REQ-001/generate").status_code == 202
    assert c.app.state.manager.wait_idle(60)
    g = c.get("/api/requirements/REQ-001").json()["generation"]
    assert g["state"] == "done" and g["backend"] == "mock"


# ------------------------------------------------------------------------------------------------------------------------ decisions
def test_decisions_history_and_summary(client):
    analyze(client)
    before = client.get("/api/projects/1/summary").json()["open_findings"]
    r = client.post("/api/requirements/REQ-003/decision", json={"decision": "accepted", "note": "ok for now"})
    assert r.status_code == 200 and r.json()["decision"] == "accepted" and r.json()["note"] == "ok for now"
    assert error_of(client.post("/api/requirements/REQ-003/decision", json={"decision": "edited"}))["code"] == "edited_text_required"
    assert client.post("/api/requirements/REQ-003/decision", json={"decision": "maybe"}).status_code == 422
    r = client.post("/api/requirements/REQ-003/decision", json={"decision": "edited", "edited_text": "The system shall respond within 2 seconds.", "note": "made measurable"})
    assert r.json()["decision"] == "edited" and [h["decision"] for h in r.json()["history"]] == ["accepted", "edited"]
    d = client.get("/api/requirements/REQ-003").json()
    assert d["decision"]["edited_text"] == "The system shall respond within 2 seconds." and len(d["decision_history"]) == 2
    assert d["text"].startswith("The system shall be easy to use")  # the original text is kept
    assert reqs(client)["items"][2]["decision"] == "edited"
    after = client.get("/api/projects/1/summary").json()
    assert after["open_findings"] < before and after["decided"] == 1


def test_reanalysis_keeps_ids_and_decisions(client):
    analyze(client)
    client.post("/api/requirements/REQ-002/decision", json={"decision": "rejected"})
    upload(client, [("more.txt", b"4.1 The system shall print a receipt for every payment.\n", "text/plain")])
    client.post("/api/projects/1/analyze")
    assert client.app.state.manager.wait_idle(120)
    ids = [i["id"] for i in reqs(client)["items"]]
    assert ids[:3] == ["REQ-001", "REQ-002", "REQ-003"] and ids[3] == "REQ-004"
    assert client.get("/api/requirements/REQ-002").json()["decision"]["decision"] == "rejected"


def test_requirement_ids_are_per_project(client):
    analyze(client)
    p2 = client.post("/api/projects", json={"name": "Other"}).json()["id"]
    analyze(client, pid=p2)
    r = client.get("/api/requirements/REQ-001")
    assert r.status_code == 409 and error_of(r)["code"] == "ambiguous_requirement_id"
    assert client.get("/api/requirements/p1-REQ-001").json()["project_id"] == 1
    assert client.get(f"/api/requirements/p{p2}-REQ-001").json()["project_id"] == p2


# ------------------------------------------------------------------------------------------------------------------------- export
def test_export_csv_and_json(client):
    analyze(client)
    client.post("/api/requirements/REQ-003/decision", json={"decision": "edited", "edited_text": "=The system shall respond within 2 seconds.", "note": "n"})
    r = client.get("/api/projects/1/export", params={"format": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv") and "attachment" in r.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [x["id"] for x in rows] == ["REQ-001", "REQ-002", "REQ-003"]
    assert rows[2]["decision"] == "edited" and rows[2]["final_text"].startswith("'=") and rows[2]["original_text"].startswith("The system shall be easy")
    assert rows[0]["duplicate_of"] == "REQ-002" and rows[0]["generation_state"] == "done" and json.loads(rows[0]["acceptance_criteria"])

    j = client.get("/api/projects/1/export", params={"format": "json"}).json()
    assert j["schema_version"] == "c1.clarified_requirements.v1" and j["project"] == {"id": 1, "name": "Core Banking Platform"}
    assert j["source_files"] == ["sample_srs.txt"] and len(j["requirements"]) == 3
    rec = j["requirements"][2]
    assert rec["decision"]["status"] == "edited" and rec["final_text"].startswith("=The system") and rec["original_text"].startswith("The system shall be easy")
    assert rec["source"] == {"file": "sample_srs.txt", "page": None, "paragraph": 2}
    assert rec["quality"]["weak_words"] == ["appropriate"] and rec["draft"]["status"] == "ok"
    assert j["requirements"][0]["decision"] is None and j["requirements"][0]["quality"]["duplicates"] == ["REQ-002"]
    assert error_of(client.get("/api/projects/1/export", params={"format": "xml"}))["code"] == "invalid_parameter"


def test_cors_allows_the_dev_ui_origin_only(client):
    ok = client.get("/api/health", headers={"Origin": "http://localhost:5173"})
    assert ok.headers.get("access-control-allow-origin") == "http://localhost:5173"
    assert client.get("/api/health", headers={"Origin": "http://127.0.0.1:5173"}).headers.get("access-control-allow-origin") == "http://127.0.0.1:5173"
    assert "access-control-allow-origin" not in client.get("/api/health", headers={"Origin": "http://evil.example"}).headers
